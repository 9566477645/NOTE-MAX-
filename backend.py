import os, re
from io import BytesIO
from gtts import gTTS
from fpdf import FPDF
from pptx import Presentation
from youtube_transcript_api import YouTubeTranscriptApi
from llama_index.core import VectorStoreIndex, SimpleDirectoryReader, Document
from llama_index.llms.ollama import Ollama
from llama_index.embeddings.huggingface import HuggingFaceEmbedding
from llama_index.core.memory import ChatMemoryBuffer
import speech_recognition as sr
from youtube_transcript_api import YouTubeTranscriptApi
class NeuralSyncBackend:
    def __init__(self):
        self.languages = {"English": "en", "Tamil": "ta", "Hindi": "hi", "Malayalam": "ml"}

    # Change this line in backend.py:
    @staticmethod
    def load_models(model_name="llama3.2"):  # 'model_name' parameter-a add pannunga
        llm = Ollama(model=model_name, request_timeout=600.0)
        embed = HuggingFaceEmbedding(model_name="BAAI/bge-small-en-v1.5")
        return llm, embed

    @staticmethod
    def fetch_yt_transcript(url):
        try:
        # Extract video ID
            v_id_match = re.search(r"(?:v=|\/)([0-9A-Za-z_-]{11}).*", url)
            if not v_id_match:
                return "Error: Invalid URL."

            v_id = v_id_match.group(1)

        # NEW METHOD (works with latest version)
            ytt_api = YouTubeTranscriptApi()
            transcript = ytt_api.fetch(v_id)

            return " ".join([t.text for t in transcript])

        except Exception as e:
            return f"Error: {str(e)}"

    @staticmethod
    def process_voice_to_text(audio_data):
        r = sr.Recognizer()
        try:
            with BytesIO(audio_data['bytes']) as f:
                with sr.AudioFile(f) as source:
                    audio = r.record(source)
                    return r.recognize_google(audio)
        except: return None

    @staticmethod
    def create_pdf_report(content, title="NeuralSync Report"):
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", 'B', 16)
        pdf.cell(200, 10, txt=title, ln=True, align='C')
        pdf.ln(10)
        pdf.set_font("Arial", size=12)
        pdf.multi_cell(0, 10, txt=content.encode('latin-1', 'ignore').decode('latin-1'))
        return pdf.output(dest='S').encode('latin-1')

    @staticmethod
    def create_pptx(slide_content):
        prs = Presentation()
        for entry in slide_content.split("Slide"):
            if ":" in entry:
                title_part, content_part = entry.split(":", 1)
                slide = prs.slides.add_slide(prs.slide_layouts[1])
                slide.shapes.title.text = title_part.strip()
                slide.placeholders[1].text = content_part.strip()
        binary_output = BytesIO()
        prs.save(binary_output)
        return binary_output.getvalue()

    @staticmethod
    def parse_quiz(text):
        questions = []
        for p in re.split(r"(?i)QUESTION:", text)[1:]:
            lines = [line.strip() for line in p.strip().split('\n') if line.strip()]
            if len(lines) < 2: continue
            q_text = lines[0]
            opts = [l for l in lines if re.match(r"^[A-D][\)\.]", l)]
            ans = re.search(r"(?i)ANSWER:\s*([A-D])", p)
            if q_text and len(opts) >= 2 and ans:
                questions.append({"q": q_text, "opts": opts, "correct": ans.group(1).upper()})
        return questions

    @staticmethod
    def extract_important_qs(text):
        return [l.strip() for l in text.split('\n') if len(l.strip()) > 10 and '?' in l]
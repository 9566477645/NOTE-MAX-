import streamlit as st
import os, re
from gtts import gTTS
from fpdf import FPDF
from pptx import Presentation 
from youtube_transcript_api import YouTubeTranscriptApi
from llama_index.core import VectorStoreIndex, SimpleDirectoryReader, Document
from llama_index.llms.ollama import Ollama
from llama_index.embeddings.huggingface import HuggingFaceEmbedding
from io import BytesIO

# --- Page Setup ---
st.set_page_config(page_title="NeuralSync Pro", layout="wide", page_icon="🧠")

# --- UI Styling ---
st.markdown("""
    <style>
    .stApp { background-color: #f8f9fa; }
    .transcript-card {
        background-color: #ffffff; padding: 20px; border-radius: 10px;
        border-left: 5px solid #ff4b4b; margin-top: 10px;
        box-shadow: 0 2px 8px rgba(0,0,0,0.1);
    }
    .quiz-box {
        background-color: #ffffff; padding: 20px; border-radius: 12px;
        border-left: 6px solid #00c6ff; margin-bottom: 20px;
        box-shadow: 0 4px 10px rgba(0,0,0,0.05);
    }
    .stButton>button { width: 100%; border-radius: 8px; font-weight: bold; }
    </style>
    """, unsafe_allow_html=True)

# --- Models ---
@st.cache_resource
def load_llm():
    llm = Ollama(model="llama3.2:1b", request_timeout=600.0, context_window=2048)
    embed = HuggingFaceEmbedding(model_name="BAAI/bge-small-en-v1.5")
    return llm, embed

llm, embed_model = load_llm()

# --- Helper Functions ---

def create_pdf_report(content, title="NeuralSync Report"):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(200, 10, txt=title, ln=True, align='C')
    pdf.ln(10)
    pdf.set_font("Arial", size=12)
    clean_text = content.encode('latin-1', 'ignore').decode('latin-1')
    pdf.multi_cell(0, 10, txt=clean_text)
    return pdf.output(dest='S').encode('latin-1')

def extract_important_qs(text):
    lines = text.split('\n')
    questions = [l.strip() for l in lines if len(l.strip()) > 10 and '?' in l]
    return questions

def create_pptx(slide_content):
    prs = Presentation()
    slides_data = slide_content.split("Slide")
    for entry in slides_data:
        if ":" in entry:
            title_part, content_part = entry.split(":", 1)
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = title_part.strip()
            slide.placeholders[1].text = content_part.strip()
    binary_output = BytesIO()
    prs.save(binary_output)
    return binary_output.getvalue()

def parse_quiz(text):
    questions = []
    parts = re.split(r"(?i)QUESTION:", text)
    for p in parts[1:]:
        lines = [line.strip() for line in p.strip().split('\n') if line.strip()]
        if len(lines) < 2: continue
        q_text = lines[0]
        opts = [l for l in lines if re.match(r"^[A-D][\)\.]", l)]
        ans = re.search(r"(?i)ANSWER:\s*([A-D])", p)
        if q_text and len(opts) >= 2 and ans:
            questions.append({"q": q_text, "opts": opts, "correct": ans.group(1).upper()})
    return questions

from youtube_transcript_api import YouTubeTranscriptApi

def fetch_yt_transcript(url):
    try:
        # Extract video ID
        v_id_match = re.search(r"(?:v=|\/)([0-9A-Za-z_-]{11}).*", url)
        if not v_id_match:
            return "Error: Invalid URL."

        v_id = v_id_match.group(1)

        # NEW METHOD (works with latest version)
        ytt_api = YouTubeTranscriptApi()
        transcript = ytt_api.fetch(v_id)

        return " ".join([t.text for t in transcript])

    except Exception as e:
        return f"Error: {str(e)}"

# --- States Initialization ---
for key in ["messages", "transcript", "index", "yt_docs", "query_engine", "source_type", "user_notes", "quiz_list", "important_qs"]:
    if key not in st.session_state:
        st.session_state[key] = "" if key == "user_notes" else ([] if key in ["messages", "yt_docs", "quiz_list", "important_qs"] else None)

# --- Layout ---
col_src, col_chat, col_std = st.columns([1, 1.8, 1.2])

# --- 1. SOURCES (Left) ---
with col_src:
    st.markdown("###  Knowledge Base")
    yt_disabled = (st.session_state.source_type == "pdf")
    with st.expander(" YouTube Link", expanded=(not yt_disabled)):
        yt_input = st.text_input("Paste Link", disabled=yt_disabled)
        if st.button("Import Video", disabled=yt_disabled) and yt_input:
            text = fetch_yt_transcript(yt_input)
            if "Error" not in text:
                st.session_state.yt_docs = [Document(text=text)]
                st.session_state.source_type = "youtube"
                st.success("Synced!")
            else: st.error(text)

    pdf_disabled = (st.session_state.source_type == "youtube")
    st.markdown("#### PDF Upload")
    files = st.file_uploader("Upload Files", accept_multiple_files=True, disabled=pdf_disabled)
    if files: st.session_state.source_type = "pdf"

    if st.button(" Sync Neural Engine", type="primary"):
        with st.spinner("Processing Data..."):
            if not os.path.exists("./data"): os.makedirs("./data")
            if files:
                for f in files:
                    with open(os.path.join("./data", f.name), "wb") as b: b.write(f.getbuffer())
                docs = SimpleDirectoryReader("./data").load_data()
            else: docs = st.session_state.yt_docs
            
            if docs:
                st.session_state.index = VectorStoreIndex.from_documents(docs, embed_model=embed_model)
                st.session_state.query_engine = st.session_state.index.as_query_engine(llm=llm)
                st.success("AI Ready!")
            else: st.error("Add a source first!")

# --- 2. CHAT & QUIZ (Middle) ---
with col_chat:
    tab_chat, tab_quiz, tab_imp = st.tabs([" Assistant", " Smart Quiz", " Important Qs"])    
    
    with tab_chat:
        for msg in st.session_state.messages:
            with st.chat_message(msg["role"]): st.markdown(msg["content"])
        if prompt := st.chat_input("Ask anything..."):
            st.session_state.messages.append({"role": "user", "content": prompt})
            with st.chat_message("user"): st.markdown(prompt)
            if st.session_state.query_engine:
                res = st.session_state.query_engine.query(prompt)
                st.session_state.messages.append({"role": "assistant", "content": res.response})
                st.rerun()

    with tab_quiz:
        st.markdown("### Interactive Quiz Hub")
        if st.button(" Generate 5 MCQs"):
            if st.session_state.query_engine:
                with st.spinner("AI is analyzing content..."):
                    q_prompt = "(important)**Generate 5 short MCQs. Strict format: QUESTION: (text) A) (opt) B) (opt) C) (opt) D) (opt) ANSWER: (Letter)"
                    res = st.session_state.query_engine.query(q_prompt)
                    st.session_state.quiz_list = parse_quiz(res.response)
            else: st.warning("Sync source first!")
        
        # Displaying Quiz
        for i, q in enumerate(st.session_state.quiz_list):
            st.markdown(f'<div class="quiz-box"><b>Q{i+1}: {q["q"]}</b></div>', unsafe_allow_html=True)
            choice = st.radio(f"Select Answer for Q{i+1}:", q["opts"], key=f"q_{i}")
            if st.button(f"Check Q{i+1}", key=f"btn_{i}"):
                if choice.strip().upper().startswith(q["correct"]): st.success(f"Correct! 🎉")
                else: st.error(f"Incorrect. Correct answer is {q['correct']}")

    with tab_imp: 
        st.markdown("###  AI Exam Predictor")
        if st.button(" Auto-Generate Key Questions"):
            if st.session_state.query_engine:
                with st.spinner("Analyzing priorities..."):
                    res = st.session_state.query_engine.query("Identify 5 high-priority conceptual questions from this content for an exam. Provide only questions.")
                    st.session_state.important_qs = extract_important_qs(str(res))
            else: st.warning("First sync pannunga macha!")
    
        for q in st.session_state.important_qs:
            st.markdown(f'<div style="background-color: #fff4e5; padding: 15px; border-radius: 10px; border-left: 5px solid #ffa500; margin-bottom: 10px;">💡 {q}</div>', unsafe_allow_html=True)

# --- 3. STUDIO & EXPORTS (Right) ---
with col_std:
    st.markdown("###  Studio Settings")
    languages = {"English": "en", "Tamil": "ta", "Hindi": "hi", "Malayalam": "ml"}
    selected_lang_name = st.selectbox("Language:", list(languages.keys()))
    
    if st.button(" Reset All"):
        st.session_state.quiz_list = []; st.session_state.messages = []; st.session_state.important_qs = []; st.rerun()

    st.divider()
    st.markdown("###  Exports")
    if st.button(" Generate Report"):
        if st.session_state.query_engine:
            with st.spinner("Drafting..."):
                res = st.session_state.query_engine.query("Write a detailed professional report.")
                st.download_button(" Download PDF", create_pdf_report(res.response), "Source_Report.pdf")

    if st.button(" Export Important Qs (PDF)"):
        if st.session_state.important_qs:
            full_text = "IMPORTANT QUESTIONS:\n\n" + "\n".join(st.session_state.important_qs)
            st.download_button(" Download Qs PDF", create_pdf_report(full_text, "Exam Priority Questions"), "Important_Qs.pdf")
        else: st.warning("Generate questions first!")

    if st.button(" Generate Slides"):
        if st.session_state.query_engine:
            with st.spinner("Designing..."):
                res = st.session_state.query_engine.query("Create 5 slides. Format: 'Slide Title: Slide Content'")
                st.download_button(" Download PPTX", create_pptx(res.response), "Sidedeck.pptx")

    st.divider()
    st.session_state.user_notes = st.text_area("Draft ideas:", value=st.session_state.user_notes, height=100)

    if st.button(f" Generate {selected_lang_name} Audio"):
        if st.session_state.query_engine:
            with st.spinner("Generating..."):
                res = st.session_state.query_engine.query(f"Summarize in {selected_lang_name}.")
                st.session_state.transcript = res.response
                gTTS(text=res.response, lang=languages[selected_lang_name]).save("audio_output.mp3")
                st.rerun()

    if st.session_state.transcript:
        st.markdown(f'<div class="transcript-card">{st.session_state.transcript}</div>', unsafe_allow_html=True)
    if os.path.exists("audio_output.mp3"): st.audio("audio_output.mp3")